import copy
from pptx import Presentation
from pptx.util import Inches

SRC = "presentation_format.pptx"
OUT = "docs/Route-Resilience-Idea-Submission.pptx"
IMG_ARCH = "docs/Prob4Extensivediagram.png"
IMG_FLOW = "docs/Prob4ShortDia.png"
BLANK = "________________"

prs = Presentation(SRC)


def set_para_text(para, text):
    runs = para.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        para.add_run().text = text


def fill_content(tf, header, bullets):
    """Keep template formatting: reuse para[0] for header, clone body para for bullets."""
    paras = tf.paragraphs
    body_ref = copy.deepcopy(paras[1]._p if len(paras) > 1 else paras[0]._p)
    set_para_text(paras[0], header)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    txBody = tf._txBody
    for b in bullets:
        newp = copy.deepcopy(body_ref)
        txBody.append(newp)
        set_para_text(tf.paragraphs[-1], b)


def append_after_label(shape, value):
    para = shape.text_frame.paragraphs[0]
    if para.runs:
        para.runs[0].text = para.runs[0].text.rstrip() + " " + value
    else:
        para.add_run().text = value


def add_fit(slide, path, top_in, maxw_in=9.0, maxh_in=3.2):
    pic = slide.shapes.add_picture(path, Inches(0.5), Inches(top_in), width=Inches(maxw_in))
    if pic.height > Inches(maxh_in):
        ratio = pic.width / pic.height
        pic.height = Inches(maxh_in)
        pic.width = int(Inches(maxh_in) * ratio)
    pic.left = int((Inches(10) - pic.width) / 2)
    pic.top = Inches(top_in)
    return pic


s = prs.slides

# ---------- Slide 1: Title ----------
for sh in s[0].shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip().lower()
    if t.startswith("team name"):
        append_after_label(sh, "Dubai_Kurukku_sandhu")
    elif t.startswith("team leader"):
        append_after_label(sh, BLANK)
    elif t.startswith("problem statement"):
        append_after_label(sh, "Route Resilience — Occlusion-Robust Road Extraction & "
                                "Graph-Theoretic Criticality Analysis for Urban Mobility")

# ---------- Slide 2: Team Members table ----------
for sh in s[1].shapes:
    if sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    txt = para.text.strip().lower()
                    if txt.startswith("name:"):
                        append_after_label_cell = para
                        if para.runs:
                            para.runs[0].text = para.runs[0].text.rstrip() + " " + BLANK
                    elif txt.startswith("college:"):
                        if para.runs:
                            para.runs[0].text = para.runs[0].text.rstrip() + " " + BLANK

# ---------- Slide 3: Opportunity ----------
fill_content(s[2].shapes[-1].text_frame, "Opportunity — what makes this different", [
    "Different: we optimise road CONNECTIVITY (clDice topology loss) and HEAL broken masks into a routable graph — not just ship pixel masks like standard extractors.",
    "Solves the problem: sees through canopy/shadow occlusions, reconnects gaps, then flags the intersections whose failure would break the city.",
    "USP: cycle-preserving graph healing + a global-efficiency Resilience Index + a plain-English LangGraph assistant for non-technical planners.",
])

# ---------- Slide 4: Features ----------
fill_content(s[3].shapes[-1].text_frame, "List of features offered by the solution", [
    "Occlusion-robust road segmentation (U-Net/SegFormer + topology-preserving clDice loss).",
    "Topological healing: skeletonize + RDP cleanup + cycle-preserving gated gap-bridging → connected, weighted, routable graph.",
    "Criticality map: betweenness centrality flags 'Gatekeeper Nodes' / bottlenecks.",
    "Node-ablation stress test → Resilience Index (global network efficiency, survives disconnection).",
    "Interactive dashboard: click a node to disable it → live reroute + travel-time increase.",
    "Conversational LangGraph/LangChain assistant (Claude) — ask resilience questions in plain English.",
])

# ---------- Slide 5: Process flow + short diagram ----------
fill_content(s[4].shapes[-1].text_frame, "Process Flow", [
    "Imagery → Segmentation → Topological Healing → Graph Analysis → Interactive Dashboard.",
])
add_fit(s[4], IMG_FLOW, top_in=2.1, maxw_in=9.0, maxh_in=3.2)

# ---------- Slide 6: Wireframe / mock (optional) ----------
fill_content(s[5].shapes[-1].text_frame, "Wireframe / Mock of the dashboard", [
    "Dark urban map; road network glows blue; key intersections pulse red, scaled by betweenness centrality.",
    "Side chat panel: user asks 'if the main highway floods, what happens?' — map greys out the closure, highlights detours, shows the Resilience Index drop.",
])

# ---------- Slide 7: Architecture diagram ----------
fill_content(s[6].shapes[-1].text_frame, "Architecture diagram of the proposed solution", [])
add_fit(s[6], IMG_ARCH, top_in=1.6, maxw_in=9.2, maxh_in=3.7)

# ---------- Slide 8: Technologies ----------
fill_content(s[7].shapes[-1].text_frame, "Technologies to be used in the solution", [
    "ML / CV: PyTorch, segmentation-models-pytorch, Albumentations.",
    "Geospatial / Graph: Rasterio, GDAL, scikit-image, sknw, NetworkX, OSMnx.",
    "App / AI: Streamlit, folium/Leaflet, LangGraph, LangChain, Anthropic Claude.",
    "Data: Cartosat-3 (primary) + pan-sharpened LISS-IV; Sentinel-2 / DeepGlobe / SpaceNet (pre-train); OSM ground truth.",
])

# ---------- Slide 9: Estimated cost (optional) ----------
fill_content(s[8].shapes[-1].text_frame, "Estimated implementation cost (optional)", [
    "Compute: 1 owned GPU workstation for fine-tuning — ~zero incremental cost.",
    "LLM API (Claude) for the assistant: a few dollars for the demo.",
    "Data: open datasets + OSM are free; Cartosat-3 provided at the event.",
    "Graph analysis + dashboard: CPU-only — no cloud cost.",
])

prs.save(OUT)
print("SAVED", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
